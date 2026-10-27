"""Bounded PDF extraction and exports from user-confirmed text only."""
import io
import json
import re
import subprocess
import sys
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from fastapi import HTTPException

BASE_DIR = Path(__file__).resolve().parents[1]
MAX_FILE = 10 * 1024 * 1024
SECTIONS = {'contacts': 'Контакты', 'summary': 'О себе', 'skills': 'Навыки',
            'experience': 'Опыт', 'projects': 'Проекты', 'education': 'Образование',
            'achievements': 'Достижения', 'languages': 'Языки'}


def extract_pdf(data):
    if len(data) > MAX_FILE:
        raise HTTPException(413, 'PDF превышает 10 МБ.')
    if not data.startswith(b'%PDF-'):
        raise HTTPException(422, 'Файл не является PDF. Выберите PDF с текстовым слоем.')
    try:
        result = subprocess.run([sys.executable, '-B', str(Path(__file__).with_name('pdf_worker.py'))],
                                input=data, capture_output=True, timeout=12, check=False)
        payload = json.loads(result.stdout)
    except (subprocess.TimeoutExpired, ValueError):
        raise HTTPException(422, 'PDF слишком сложный для обработки. Сохраните упрощённый PDF или вставьте текст.') from None
    if result.returncode or 'error' in payload:
        messages = {
            'encrypted': 'PDF защищён паролем. Загрузите незашифрованную копию.',
            'pages': 'Допускается PDF от 1 до 20 страниц.',
            'text_limit': 'В PDF слишком много текста. Оставьте только резюме.',
            'scan': 'В PDF не найден читаемый текст. OCR пока не поддерживается: вставьте текст вручную или загрузите PDF с текстовым слоем.',
        }
        raise HTTPException(422, messages.get(payload.get('error'), 'Не удалось прочитать PDF. Проверьте файл или вставьте текст.'))
    return payload['pages']


def assemble_resume(app):
    replacements = {c.fact_id: c.after for c in app.changes if c.decision == 'accepted'}
    order = list(dict.fromkeys([*app.section_order, *SECTIONS]))
    blocks = []
    for section in order:
        facts = [replacements.get(f.id, f.text) for f in app.facts if f.section == section]
        if facts:
            blocks.append('## ' + SECTIONS[section] + '\n' + '\n\n'.join(facts))
    return '\n\n'.join(blocks)


def docx_bytes(text):
    doc = Document()
    doc.core_properties.title = 'Career application'
    for line in text.splitlines():
        if line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('# '):
            doc.add_heading(line[2:], level=0)
        else:
            doc.add_paragraph(line)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def pdf_bytes(text, title='Резюме'):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    if 'Manrope' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('Manrope', str(BASE_DIR / 'web/static/fonts/Manrope-Regular.ttf')))
        pdfmetrics.registerFont(TTFont('ManropeSemibold', str(BASE_DIR / 'web/static/fonts/Manrope-Semibold.ttf')))
    body = ParagraphStyle('body', fontName='Manrope', fontSize=10, leading=16,
                          textColor=colors.HexColor('#26362f'), spaceAfter=8, splitLongWords=True)
    heading = ParagraphStyle('heading', parent=body, fontName='ManropeSemibold', fontSize=13, leading=18, spaceBefore=14,
                             spaceAfter=8, textColor=colors.HexColor('#175b43'), keepWithNext=True)
    title_style = ParagraphStyle('title', parent=heading, fontSize=20, leading=28)
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=48, leftMargin=48,
                            topMargin=40, bottomMargin=44, title=title, author='')
    story = []
    for block in re.split(r'\n\s*\n', text.strip()):
        if not block.strip():
            continue
        for line in block.splitlines():
            is_heading = line.startswith('#')
            clean = re.sub(r'^#{1,3}\s*', '', line)
            story.append(Paragraph(escape(clean), title_style if line.startswith('# ') else heading if is_heading else body))
    if not story:
        raise HTTPException(422, 'Документ пуст.')
    def footer(canvas, _doc):
        canvas.setFont('Manrope', 8)
        canvas.setFillColor(colors.HexColor('#687a70'))
        canvas.drawRightString(A4[0] - 48, 24, str(_doc.page))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
